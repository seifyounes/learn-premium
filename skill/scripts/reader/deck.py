""".pptx Materials: each slide's text and notes, its audio (transcribed) and its embedded video."""

import json
import zipfile
from pathlib import Path

from lxml import etree
from reader.errors import BadInput

AUDIO = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/audio"
VIDEO = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/video"
MEDIA = "http://schemas.microsoft.com/office/2007/relationships/media"
NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "p14": "http://schemas.microsoft.com/office/powerpoint/2010/main",
}
# PowerPoint's Record Slide Show marks the audio it records as narration in the slide's timing.
NARRATED_SHAPES = etree.XPath(
    ".//p:timing//p:audio[@isNarration='1']//p:spTgt/@spid", namespaces=NS)
MEDIA_RIDS = etree.XPath(
    ".//p:pic[p:nvPicPr/p:cNvPr/@id=$spid]/p:nvPicPr/p:nvPr//@r:link"
    " | .//p:pic[p:nvPicPr/p:cNvPr/@id=$spid]/p:nvPicPr/p:nvPr//@r:embed", namespaces=NS)


def _texts(shapes, skip_id=None):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        if shape.shape_id == skip_id:
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _texts(shape.shapes)
        elif shape.has_text_frame and shape.text_frame.text.strip():
            yield shape.text_frame.text
        elif shape.has_table:
            for row in shape.table.rows:
                yield " | ".join(cell.text for cell in row.cells)


def _media_key(relationship):
    """One media file, however many relationships link it (a clip has an audio or video one and
    a media one)."""
    if relationship.is_external:
        return relationship.target_ref
    return relationship.target_part.partname


def _media(slide):
    """The slide's media: (key, "audio" or "video", one relationship to it)."""
    found = {}
    for relationship in slide.part.rels.values():
        if relationship.reltype in (AUDIO, VIDEO, MEDIA):
            reltypes, _ = found.setdefault(_media_key(relationship), (set(), relationship))
            reltypes.add(relationship.reltype)
    for key, (reltypes, relationship) in found.items():
        if VIDEO in reltypes:
            kind = "video"
        elif AUDIO in reltypes or relationship.is_external:
            kind = "audio"
        else:
            kind = "video" if relationship.target_part.content_type.startswith("video/") else "audio"
        yield key, kind, relationship


def _narrated(slide):
    """Keys (as `_media_key` gives them) of the media PowerPoint recorded as narration."""
    relationships = slide.part.rels
    return {_media_key(relationships[rid])
            for spid in NARRATED_SHAPES(slide._element)
            for rid in MEDIA_RIDS(slide._element, spid=spid) if rid in relationships}


def _read_slide(number, slide, out, transcriber):
    title = slide.shapes.title
    notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
    entry = {"number": number, "title": title.text if title is not None else "",
             "text": "\n".join(_texts(slide.shapes, title.shape_id if title is not None else None)),
             "notes": notes, "audio": [], "video": []}
    narrated = _narrated(slide)
    for key, kind, relationship in _media(slide):
        name = f"slide-{number:02d}-{kind}-{len(entry[kind]) + 1}"
        item = {}
        if relationship.is_external:
            item["linked"] = relationship.target_ref
        else:
            (out / "media").mkdir(exist_ok=True)
            item["file"] = out / "media" / f"{name}.{relationship.target_part.partname.ext}"
            item["file"].write_bytes(relationship.target_part.blob)
        if kind == "audio":
            item["narration"] = key in narrated
            if "file" in item:
                (out / "transcripts").mkdir(exist_ok=True)
                item["transcript"] = out / "transcripts" / f"{name}.json"
                item["transcript"].write_text(
                    json.dumps(transcriber(item["file"]), indent=2, ensure_ascii=False),
                    encoding="utf-8")
        else:
            # Deck-embedded video is source only: never re-hosted, never copied into a Course project.
            item["sourceOnly"] = True
        entry[kind].append(item)
    return entry


def read_deck(material: Path, out: Path, transcriber) -> dict:
    from pptx import Presentation

    with open(material, "rb") as f:
        try:
            presentation = Presentation(f)
        except (zipfile.BadZipFile, KeyError, ValueError) as e:
            raise BadInput(f"not a .pptx deck ({e})") from None
        slides = [_read_slide(number, slide, out, transcriber)
                  for number, slide in enumerate(presentation.slides, start=1)]
    return {"kind": "deck", "slides": slides}
