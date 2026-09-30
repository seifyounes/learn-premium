"""Synthetic Materials for the Materials reader tests: PDFs written byte by byte and .pptx decks
built with python-pptx, with per-slide audio (a narration timing node, as PowerPoint's Record Slide
Show writes it) and an embedded video. No real Materials."""

import io
import math
import struct
import wave
from pathlib import Path

DRAWN = "0 0 0 rg 72 72 300 200 re f BT /F1 24 Tf 72 700 Td (Q = kA dT/dx) Tj ET"
WHITE_ON_WHITE = "1 1 1 rg BT /F1 24 Tf 72 700 Td (invisible) Tj ET"


def pdf(pages):
    """A PDF whose pages draw the given content streams (None: a page with no content at all)."""
    objects = ["<< /Type /Catalog /Pages 2 0 R >>", None,
               "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for content in pages:
        page = len(objects) + 1
        kids.append(f"{page} 0 R")
        if content is None:
            objects.append("<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] >>")
            continue
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                       f"/Resources << /Font << /F1 3 0 R >> >> /Contents {page + 1} 0 R >>")
        objects.append(f"<< /Length {len(content)} >>\nstream\n{content}\nendstream")
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for n, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{n} 0 obj\n{body}\nendobj\n".encode("latin-1"))
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
              .encode())
    return out.getvalue()


def tone(seconds=0.5, hz=440):
    """A short mono WAV: stands in for a slide's narration audio."""
    rate = 16000
    frames = b"".join(struct.pack("<h", int(8000 * math.sin(2 * math.pi * hz * i / rate)))
                      for i in range(int(rate * seconds)))
    out = io.BytesIO()
    with wave.open(out, "wb") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(rate)
        w.writeframes(frames)
    return out.getvalue()


# Not a playable film: the reader copies media bytes, it never decodes video.
VIDEO = b"\x00\x00\x00\x18ftypmp42\x00\x00\x00\x00mp42isom" + b"fixture video " * 64


def _tiny_png():
    from PIL import Image

    out = io.BytesIO()
    Image.new("RGB", (8, 8), (200, 30, 30)).save(out, "PNG")
    return out.getvalue()


def _add_narration(slide, audio, shape_id):
    """Embed `audio` on `slide` the way a recorded narration is stored: an audio shape linked
    through an audio and a media relationship to one media part, and an isNarration timing node."""
    from lxml import etree
    from pptx.media import Video
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.parts.media import MediaPart

    # python-pptx's media class takes any media blob; this one is audio.
    media = MediaPart.new(slide.part.package, Video(audio, "audio/x-wav", "narration.wav"))
    audio_rid = slide.part.relate_to(media, RT.AUDIO)
    media_rid = slide.part.relate_to(media, "http://schemas.microsoft.com/office/2007/relationships/media")
    image_rid = slide.part.get_or_add_image_part(io.BytesIO(_tiny_png()))[1]

    ns = ('xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
          'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
          'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
          'xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main"')
    pic = etree.fromstring(f"""
<p:pic {ns}>
  <p:nvPicPr>
    <p:cNvPr id="{shape_id}" name="Recorded Sound"/>
    <p:cNvPicPr/>
    <p:nvPr>
      <a:audioFile r:link="{audio_rid}"/>
      <p:extLst><p:ext uri="{{DAA4B4D4-6D71-4841-9C94-3DE7FCFB9230}}">
        <p14:media r:embed="{media_rid}"/>
      </p:ext></p:extLst>
    </p:nvPr>
  </p:nvPicPr>
  <p:blipFill><a:blip r:embed="{image_rid}"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>
  <p:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="304800" cy="304800"/></a:xfrm>
    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>
</p:pic>""")
    slide.shapes._spTree.append(pic)
    timing = etree.fromstring(f"""
<p:timing {ns}><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">
  <p:childTnLst>
    <p:audio isNarration="1"><p:cMediaNode vol="80000" showWhenStopped="0">
      <p:cTn id="2" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>
      </p:cTn>
      <p:tgtEl><p:spTgt spid="{shape_id}"/></p:tgtEl>
    </p:cMediaNode></p:audio>
  </p:childTnLst>
</p:cTn></p:par></p:tnLst></p:timing>""")
    slide._element.append(timing)


def deck(path: Path, slides):
    """Write a .pptx. Each slide is a dict: `title`, optional `body`, `notes`, `narration` (WAV
    bytes) and `video` (MP4 bytes)."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for spec in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = spec["title"]
        slide.placeholders[1].text = spec.get("body", "")
        if "notes" in spec:
            slide.notes_slide.notes_text_frame.text = spec["notes"]
        if "video" in spec:
            slide.shapes.add_movie(io.BytesIO(spec["video"]), Inches(1), Inches(1), Inches(4),
                                   Inches(3), poster_frame_image=io.BytesIO(_tiny_png()),
                                   mime_type="video/mp4")
        if "narration" in spec:
            _add_narration(slide, spec["narration"], shape_id=900)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path
